"""T-18 보고서 생성기 (DOCX / PDF)."""
from __future__ import annotations

import traceback
from datetime import datetime
from pathlib import Path
from typing import Any

from config import SESSION_DIR


def generate_report(session_id: str, analysis_result: dict[str, Any], output_format: str = "docx") -> dict:
    """분석 결과를 보고서 파일로 저장한다."""
    report_dir = SESSION_DIR / session_id / "reports"
    report_dir.mkdir(parents=True, exist_ok=True)

    ext = "pdf" if output_format == "pdf" else "md" if output_format == "md" else "docx"
    output_path = report_dir / f"analysis_report_{datetime.now().strftime('%H%M%S')}.{ext}"

    try:
        if output_format == "pdf":
            _build_pdf(output_path, analysis_result)
        elif output_format == "md":
            _build_md(output_path, analysis_result)
        else:
            _build_docx(output_path, analysis_result)
        return {"success": True, "report_path": str(output_path), "error": None}
    except Exception:
        return {"success": False, "report_path": "", "error": traceback.format_exc()}


def _build_docx(output_path: Path, data: dict[str, Any]) -> None:
    from docx import Document
    from docx.shared import Inches

    doc = Document()
    doc.add_heading("E_LENS 이커머스 분석 보고서", level=0)
    doc.add_paragraph(f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M')}")

    if summary := data.get("summary"):
        doc.add_heading("분석 요약", level=1)
        doc.add_paragraph(str(summary))

    if insights := data.get("insights"):
        doc.add_heading("핵심 인사이트", level=1)
        for ins in insights:
            doc.add_paragraph(str(ins), style="List Bullet")

    if actions := data.get("actions"):
        doc.add_heading("액션 아이템", level=1)
        for act in actions:
            doc.add_paragraph(str(act), style="List Bullet")

    if kpi := data.get("kpi_result"):
        doc.add_heading("KPI 결과", level=1)
        table = doc.add_table(rows=1, cols=2)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "지표"
        table.rows[0].cells[1].text = "값"
        for k, v in kpi.items():
            row = table.add_row().cells
            row[0].text = str(k)
            row[1].text = str(round(v, 4) if isinstance(v, float) else v)

    if ranking := data.get("feature_ranking"):
        doc.add_heading("주요 변수 (상위 5)", level=1)
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        table.rows[0].cells[0].text = "순위"
        table.rows[0].cells[1].text = "변수명"
        table.rows[0].cells[2].text = "중요도"
        for i, (feat, score) in enumerate(list(ranking.items())[:5], 1):
            row = table.add_row().cells
            row[0].text = str(i)
            row[1].text = str(feat)
            row[2].text = str(score)

    if image_paths := data.get("image_paths"):
        doc.add_heading("시각화", level=1)
        for img_path in image_paths:
            p = Path(str(img_path))
            if p.is_file() and p.suffix.lower() == ".png":
                doc.add_picture(str(p), width=Inches(5.8))

    doc.save(str(output_path))


def _build_pdf(output_path: Path, data: dict[str, Any]) -> None:
    """reportlab 기반 PDF 보고서 생성."""
    try:
        from reportlab.lib import colors
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import HRFlowable, Image, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    except ImportError as exc:
        raise ImportError("reportlab 미설치: uv add reportlab") from exc

    font_name = "Helvetica"
    for font_path in [
        "/System/Library/Fonts/AppleGothic.ttf",
        "/Library/Fonts/NanumGothic.ttf",
        "/usr/share/fonts/truetype/nanum/NanumGothic.ttf",
    ]:
        p = Path(font_path)
        if p.exists():
            try:
                pdfmetrics.registerFont(TTFont("Korean", str(p)))
                font_name = "Korean"
                break
            except Exception:
                continue

    doc = SimpleDocTemplate(
        str(output_path),
        pagesize=A4,
        leftMargin=2.5 * cm,
        rightMargin=2.5 * cm,
        topMargin=2.0 * cm,
        bottomMargin=2.0 * cm,
    )
    story: list[Any] = []
    blue = colors.HexColor("#2E74B5")
    gray = colors.HexColor("#707070")

    title_style = ParagraphStyle("title", fontName=font_name, fontSize=18, textColor=blue, alignment=1, spaceAfter=8)
    date_style = ParagraphStyle("date", fontName=font_name, fontSize=9, textColor=gray, alignment=1, spaceAfter=12)
    h1_style = ParagraphStyle("h1", fontName=font_name, fontSize=12, textColor=blue, spaceBefore=12, spaceAfter=6)
    body_style = ParagraphStyle("body", fontName=font_name, fontSize=10, leading=15, spaceAfter=4)
    bullet_style = ParagraphStyle("bullet", fontName=font_name, fontSize=10, leading=15, leftIndent=12)

    story.append(Paragraph("E_LENS 이커머스 분석 보고서", title_style))
    story.append(Paragraph(f"생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M')}", date_style))
    story.append(HRFlowable(width="100%", thickness=1, color=blue))
    story.append(Spacer(1, 10))

    if summary := data.get("summary"):
        story.append(Paragraph("분석 요약", h1_style))
        story.append(Paragraph(str(summary), body_style))

    if insights := data.get("insights"):
        story.append(Paragraph("핵심 인사이트", h1_style))
        for ins in insights:
            story.append(Paragraph(f"• {ins}", bullet_style))

    if actions := data.get("actions"):
        story.append(Paragraph("액션 아이템", h1_style))
        for act in actions:
            story.append(Paragraph(f"• {act}", bullet_style))

    if kpi := data.get("kpi_result"):
        story.append(Paragraph("KPI 결과", h1_style))
        rows = [["지표", "값"]]
        for k, v in kpi.items():
            rows.append([str(k), str(round(v, 4) if isinstance(v, float) else v)])
        table = Table(rows, colWidths=[8 * cm, 5 * cm])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), blue),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ]
            )
        )
        story.append(table)

    if ranking := data.get("feature_ranking"):
        story.append(Paragraph("주요 변수 (상위 5)", h1_style))
        rows = [["순위", "변수명", "중요도"]]
        for i, (feat, score) in enumerate(list(ranking.items())[:5], 1):
            rows.append([str(i), str(feat), str(score)])
        table = Table(rows, colWidths=[2 * cm, 8 * cm, 3 * cm])
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), blue),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, -1), font_name),
                    ("FONTSIZE", (0, 0), (-1, -1), 9),
                    ("GRID", (0, 0), (-1, -1), 0.4, colors.grey),
                ]
            )
        )
        story.append(table)

    if image_paths := data.get("image_paths"):
        story.append(Paragraph("시각화", h1_style))
        for img_path in image_paths:
            p = Path(str(img_path))
            if p.is_file() and p.suffix.lower() == ".png":
                story.append(Image(str(p), width=14 * cm, height=8.5 * cm))
                story.append(Spacer(1, 6))

    doc.build(story)


def _build_md(output_path: Path, data: dict[str, Any]) -> None:
    """Markdown 보고서 생성."""
    lines: list[str] = []
    lines.append("# E_LENS 이커머스 분석 보고서")
    lines.append(f"- 생성일시: {datetime.now().strftime('%Y-%m-%d %H:%M')}")
    lines.append("")

    if summary := data.get("summary"):
        lines.append("## 분석 요약")
        lines.append(str(summary))
        lines.append("")

    if insights := data.get("insights"):
        lines.append("## 핵심 인사이트")
        for ins in insights:
            lines.append(f"- {ins}")
        lines.append("")

    if actions := data.get("actions"):
        lines.append("## 액션 아이템")
        for act in actions:
            lines.append(f"- {act}")
        lines.append("")

    if kpi := data.get("kpi_result"):
        lines.append("## KPI 결과")
        for k, v in kpi.items():
            lines.append(f"- {k}: {round(v, 4) if isinstance(v, float) else v}")
        lines.append("")

    if ranking := data.get("feature_ranking"):
        lines.append("## 주요 변수 (상위 5)")
        for i, (feat, score) in enumerate(list(ranking.items())[:5], 1):
            lines.append(f"{i}. {feat}: {score}")
        lines.append("")

    if image_paths := data.get("image_paths"):
        lines.append("## 시각화")
        for img_path in image_paths:
            p = Path(str(img_path))
            if p.is_file():
                lines.append(f"- {p.name}")
        lines.append("")

    output_path.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")